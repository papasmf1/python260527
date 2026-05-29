import sqlite3
import sys
from datetime import datetime
from typing import List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QApplication,
    QFormLayout,
    QGridLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
    QFileDialog,
)

DB_NAME = "Bycle.db"
TABLE_NAME = "Bycle"


class BycleManagerWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Bycle Manager (SQLite3 + PyQt6)")
        self.resize(960, 660)

        self.conn = sqlite3.connect(DB_NAME)
        self._create_table()
        self._seed_data_if_empty(100)

        self._build_ui()
        self._apply_styles()
        self.load_all_data()

    def _build_ui(self) -> None:
        central = QWidget()
        central.setObjectName("central")
        self.setCentralWidget(central)

        root_layout = QVBoxLayout(central)
        root_layout.setContentsMargins(22, 22, 22, 22)
        root_layout.setSpacing(14)

        title = QLabel("Bycle 관리 프로그램")
        title.setObjectName("title")
        root_layout.addWidget(title)

        form_layout = QFormLayout()
        self.id_input = QLineEdit()
        self.name_input = QLineEdit()
        self.price_input = QLineEdit()
        self.qty_input = QLineEdit()

        self.id_input.setPlaceholderText("정수 ID")
        self.name_input.setPlaceholderText("자전거 이름")
        self.price_input.setPlaceholderText("가격(정수)")
        self.qty_input.setPlaceholderText("수량(정수)")

        self.id_input.setObjectName("idInput")
        self.name_input.setObjectName("nameInput")
        self.price_input.setObjectName("priceInput")
        self.qty_input.setObjectName("qtyInput")

        form_layout.addRow("ID", self.id_input)
        form_layout.addRow("이름", self.name_input)
        form_layout.addRow("가격", self.price_input)
        form_layout.addRow("수량", self.qty_input)
        root_layout.addLayout(form_layout)

        button_grid = QGridLayout()
        self.add_btn = QPushButton("입력")
        self.update_btn = QPushButton("수정")
        self.delete_btn = QPushButton("삭제")
        self.clear_btn = QPushButton("입력창 초기화")

        self.add_btn.setProperty("role", "accent")
        self.update_btn.setProperty("role", "primary")
        self.delete_btn.setProperty("role", "danger")
        self.clear_btn.setProperty("role", "ghost")

        self.add_btn.clicked.connect(self.add_record)
        self.update_btn.clicked.connect(self.update_record)
        self.delete_btn.clicked.connect(self.delete_record)
        self.clear_btn.clicked.connect(self.clear_inputs)

        self.excel_btn = QPushButton("엑셀로 출력")
        self.excel_btn.setProperty("role", "excel")
        self.excel_btn.clicked.connect(self.export_to_excel)

        button_grid.addWidget(self.add_btn, 0, 0)
        button_grid.addWidget(self.update_btn, 0, 1)
        button_grid.addWidget(self.delete_btn, 0, 2)
        button_grid.addWidget(self.clear_btn, 0, 3)
        button_grid.addWidget(self.excel_btn, 0, 4)
        root_layout.addLayout(button_grid)

        search_layout = QHBoxLayout()
        self.search_id_input = QLineEdit()
        self.search_name_input = QLineEdit()
        self.search_btn = QPushButton("검색")
        self.show_all_btn = QPushButton("전체 조회")

        self.search_id_input.setPlaceholderText("ID 검색")
        self.search_name_input.setPlaceholderText("이름 검색")

        self.search_btn.setProperty("role", "primary")
        self.show_all_btn.setProperty("role", "ghost")

        self.search_btn.clicked.connect(self.search_records)
        self.show_all_btn.clicked.connect(self.load_all_data)

        search_layout.addWidget(QLabel("검색 조건:"))
        search_layout.addWidget(self.search_id_input)
        search_layout.addWidget(self.search_name_input)
        search_layout.addWidget(self.search_btn)
        search_layout.addWidget(self.show_all_btn)
        root_layout.addLayout(search_layout)

        self.table = QTableWidget()
        self.table.setColumnCount(4)
        self.table.setHorizontalHeaderLabels(["id", "name", "price", "qty"])
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.doubleClicked.connect(self.load_selected_row_to_inputs)
        self.table.horizontalHeader().setStretchLastSection(True)
        root_layout.addWidget(self.table)

    def _apply_styles(self) -> None:
        self.setStyleSheet(
            """
            QMainWindow {
                background-color: #f5f7fb;
            }

            QWidget#central {
                background: qlineargradient(
                    x1: 0,
                    y1: 0,
                    x2: 1,
                    y2: 1,
                    stop: 0 #f8fafc,
                    stop: 1 #e7eef8
                );
            }

            QLabel {
                color: #1f2937;
                font-size: 13px;
                font-family: "Malgun Gothic", "Segoe UI";
            }

            QLabel#title {
                font-size: 30px;
                font-weight: 800;
                color: #0f172a;
                padding: 8px 4px 10px 4px;
            }

            QLineEdit {
                background-color: #ffffff;
                border: 1px solid #cbd5e1;
                border-radius: 10px;
                padding: 8px 10px;
                min-height: 22px;
                font-size: 13px;
                color: #0f172a;
                selection-background-color: #2563eb;
            }

            QLineEdit:focus {
                border: 2px solid #2563eb;
                background-color: #f8fbff;
            }

            QPushButton {
                border: none;
                border-radius: 10px;
                min-height: 34px;
                padding: 6px 14px;
                font-size: 13px;
                font-weight: 700;
                font-family: "Malgun Gothic", "Segoe UI";
            }

            QPushButton[role="accent"] {
                color: #ffffff;
                background: qlineargradient(
                    x1: 0,
                    y1: 0,
                    x2: 1,
                    y2: 0,
                    stop: 0 #0284c7,
                    stop: 1 #0ea5e9
                );
            }

            QPushButton[role="primary"] {
                color: #ffffff;
                background-color: #2563eb;
            }

            QPushButton[role="danger"] {
                color: #ffffff;
                background-color: #dc2626;
            }

            QPushButton[role="ghost"] {
                color: #1e293b;
                background-color: #e2e8f0;
            }

            QPushButton:hover {
                margin-top: -1px;
            }

            QPushButton[role="accent"]:hover {
                background: qlineargradient(
                    x1: 0,
                    y1: 0,
                    x2: 1,
                    y2: 0,
                    stop: 0 #0369a1,
                    stop: 1 #0284c7
                );
            }

            QPushButton[role="primary"]:hover {
                background-color: #1d4ed8;
            }

            QPushButton[role="danger"]:hover {
                background-color: #b91c1c;
            }

            QPushButton[role="ghost"]:hover {
                background-color: #cbd5e1;
            }

            QPushButton[role="excel"] {
                color: #ffffff;
                background-color: #16a34a;
            }

            QPushButton[role="excel"]:hover {
                background-color: #15803d;
            }

            QPushButton:pressed {
                margin-top: 1px;
            }

            QTableWidget {
                background: #ffffff;
                border: 1px solid #cbd5e1;
                border-radius: 12px;
                gridline-color: #e2e8f0;
                font-size: 13px;
                alternate-background-color: #f8fafc;
                selection-background-color: #dbeafe;
                selection-color: #0f172a;
            }

            QHeaderView::section {
                background-color: #0f172a;
                color: #f8fafc;
                border: none;
                padding: 8px;
                font-size: 12px;
                font-weight: 700;
            }
            """
        )

    def _create_table(self) -> None:
        query = f"""
        CREATE TABLE IF NOT EXISTS {TABLE_NAME} (
            id INTEGER PRIMARY KEY,
            name TEXT NOT NULL,
            price INTEGER NOT NULL,
            qty INTEGER NOT NULL
        )
        """
        self.conn.execute(query)
        self.conn.commit()

    def _seed_data_if_empty(self, count: int) -> None:
        cursor = self.conn.execute(f"SELECT COUNT(*) FROM {TABLE_NAME}")
        row_count = int(cursor.fetchone()[0])

        if row_count > 0:
            return

        sample_rows: List[Tuple[int, str, int, int]] = []
        for i in range(1, count + 1):
            sample_rows.append(
                (
                    i,
                    f"Bicycle-{i:03d}",
                    100000 + (i * 2500),
                    (i % 15) + 1,
                )
            )

        self.conn.executemany(
            f"INSERT INTO {TABLE_NAME} (id, name, price, qty) VALUES (?, ?, ?, ?)",
            sample_rows,
        )
        self.conn.commit()

    def _rows_to_table(self, rows: List[Tuple[int, str, int, int]]) -> None:
        self.table.setRowCount(len(rows))

        for r_idx, row in enumerate(rows):
            for c_idx, value in enumerate(row):
                item = QTableWidgetItem(str(value))
                if c_idx in (0, 2, 3):
                    item.setTextAlignment(
                        Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter
                    )
                self.table.setItem(r_idx, c_idx, item)

        self.table.resizeColumnsToContents()

    def _require_int(self, text: str, field_name: str) -> int:
        try:
            return int(text)
        except ValueError as exc:
            raise ValueError(f"{field_name}는 정수여야 합니다.") from exc

    def _fetch_rows(self, where_clause: str = "", params: Tuple = ()) -> List[Tuple[int, str, int, int]]:
        query = f"SELECT id, name, price, qty FROM {TABLE_NAME}"
        if where_clause:
            query += f" WHERE {where_clause}"
        query += " ORDER BY id"

        cursor = self.conn.execute(query, params)
        return cursor.fetchall()

    def load_all_data(self) -> None:
        rows = self._fetch_rows()
        self._rows_to_table(rows)

    def search_records(self) -> None:
        raw_id = self.search_id_input.text().strip()
        raw_name = self.search_name_input.text().strip()

        conditions = []
        params = []

        if raw_id:
            try:
                conditions.append("id = ?")
                params.append(int(raw_id))
            except ValueError:
                QMessageBox.warning(self, "입력 오류", "검색 ID는 정수로 입력하세요.")
                return

        if raw_name:
            conditions.append("name LIKE ?")
            params.append(f"%{raw_name}%")

        where_clause = " AND ".join(conditions)
        rows = self._fetch_rows(where_clause, tuple(params))
        self._rows_to_table(rows)

    def add_record(self) -> None:
        try:
            row_id = self._require_int(self.id_input.text().strip(), "ID")
            name = self.name_input.text().strip()
            price = self._require_int(self.price_input.text().strip(), "가격")
            qty = self._require_int(self.qty_input.text().strip(), "수량")

            if not name:
                raise ValueError("이름은 비어 있을 수 없습니다.")

            self.conn.execute(
                f"INSERT INTO {TABLE_NAME} (id, name, price, qty) VALUES (?, ?, ?, ?)",
                (row_id, name, price, qty),
            )
            self.conn.commit()
            self.load_all_data()
            QMessageBox.information(self, "완료", "데이터를 입력했습니다.")
        except sqlite3.IntegrityError:
            QMessageBox.warning(self, "입력 오류", "이미 존재하는 ID입니다.")
        except ValueError as exc:
            QMessageBox.warning(self, "입력 오류", str(exc))

    def update_record(self) -> None:
        try:
            row_id = self._require_int(self.id_input.text().strip(), "ID")
            name = self.name_input.text().strip()
            price = self._require_int(self.price_input.text().strip(), "가격")
            qty = self._require_int(self.qty_input.text().strip(), "수량")

            if not name:
                raise ValueError("이름은 비어 있을 수 없습니다.")

            cursor = self.conn.execute(
                f"UPDATE {TABLE_NAME} SET name = ?, price = ?, qty = ? WHERE id = ?",
                (name, price, qty, row_id),
            )
            self.conn.commit()

            if cursor.rowcount == 0:
                QMessageBox.warning(self, "수정 실패", "해당 ID가 없습니다.")
                return

            self.load_all_data()
            QMessageBox.information(self, "완료", "데이터를 수정했습니다.")
        except ValueError as exc:
            QMessageBox.warning(self, "입력 오류", str(exc))

    def delete_record(self) -> None:
        try:
            row_id = self._require_int(self.id_input.text().strip(), "ID")

            cursor = self.conn.execute(
                f"DELETE FROM {TABLE_NAME} WHERE id = ?",
                (row_id,),
            )
            self.conn.commit()

            if cursor.rowcount == 0:
                QMessageBox.warning(self, "삭제 실패", "해당 ID가 없습니다.")
                return

            self.load_all_data()
            QMessageBox.information(self, "완료", "데이터를 삭제했습니다.")
        except ValueError as exc:
            QMessageBox.warning(self, "입력 오류", str(exc))

    def load_selected_row_to_inputs(self) -> None:
        row = self.table.currentRow()
        if row < 0:
            return

        self.id_input.setText(self.table.item(row, 0).text())
        self.name_input.setText(self.table.item(row, 1).text())
        self.price_input.setText(self.table.item(row, 2).text())
        self.qty_input.setText(self.table.item(row, 3).text())

    def clear_inputs(self) -> None:
        self.id_input.clear()
        self.name_input.clear()
        self.price_input.clear()
        self.qty_input.clear()

    def export_to_excel(self) -> None:
        row_count = self.table.rowCount()
        if row_count == 0:
            QMessageBox.warning(self, "엑셀 출력", "내보낼 데이터가 없습니다.")
            return

        default_name = f"Bycle_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "엑셀 파일 저장",
            default_name,
            "Excel Files (*.xlsx)",
        )
        if not file_path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bycle"

        headers = ["id", "name", "price", "qty"]
        header_fill = PatternFill(fill_type="solid", fgColor="0F172A")
        header_font = Font(name="Malgun Gothic", bold=True, color="F8FAFC", size=11)
        header_align = Alignment(horizontal="center", vertical="center")

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        ws.row_dimensions[1].height = 22

        for r in range(row_count):
            for c in range(4):
                raw = self.table.item(r, c).text().strip()
                value = int(raw) if c in (0, 2, 3) else raw
                cell = ws.cell(row=r + 2, column=c + 1, value=value)
                cell.alignment = Alignment(
                    horizontal="right" if c in (0, 2, 3) else "left",
                    vertical="center",
                )
                if r % 2 == 1:
                    cell.fill = PatternFill(fill_type="solid", fgColor="F8FAFC")

        col_widths = [10, 28, 14, 10]
        for col_idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(col_idx)
            ].width = width

        try:
            wb.save(file_path)
            QMessageBox.information(
                self, "완료", f"엑셀 파일을 저장했습니다.\n{file_path}"
            )
        except PermissionError:
            QMessageBox.critical(
                self, "저장 실패", "파일이 열려 있어 저장할 수 없습니다.\n파일을 닫고 다시 시도하세요."
            )

    def closeEvent(self, event) -> None:  # type: ignore[override]
        self.conn.close()
        super().closeEvent(event)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = BycleManagerWindow()
    window.show()
    sys.exit(app.exec())
